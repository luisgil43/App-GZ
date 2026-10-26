"""Exportación filtrada de la asignación mensual."""

import csv
import io
import unicodedata
from datetime import date

from django.contrib.auth.decorators import login_required
from django.http import HttpResponse
from django.shortcuts import get_object_or_404
from django.views.decorators.http import require_GET

from planificacion.models import PlanificacionMensual, SitioPlanificado
from usuarios.decoradores import rol_requerido


COLUMNAS = [
    "ID CLARO",
    "NOMBRE",
    "REGIÓN",
    "COMUNA",
    "TIPO DE ZONA",
    "DIRECCIÓN",
    "CONTACTO / ACCESO",
    "CONDICIÓN DE ACCESO",
    "PERMISO",
    "ESTADO",
    "FECHA PLANIFICADA",
    "ALERTA",
]


def normalizar(valor):
    texto = unicodedata.normalize(
        "NFD",
        str(valor or "").strip().lower(),
    )
    return "".join(
        caracter
        for caracter in texto
        if unicodedata.category(caracter) != "Mn"
    )


def texto_seguro(valor):
    texto = str(valor or "")
    if texto.lstrip().startswith(("=", "+", "-", "@")):
        return "'" + texto
    return texto


def etiqueta_campo(objeto, campo):
    metodo = getattr(
        objeto,
        f"get_{campo}_display",
        None,
    )
    if callable(metodo):
        return str(metodo() or "")
    return str(getattr(objeto, campo, "") or "")


def obtener_valor(objeto, nombre):
    return getattr(objeto, nombre, "") or ""


def contacto_acceso(item):
    """
    Extrae información de contacto sin alterar los registros.
    """
    sitio = item.sitio

    partes = []

    for campo in (
        "nombre_contacto",
        "contacto",
        "telefono_contacto",
        "telefono",
        "correo_contacto",
        "email_contacto",
    ):
        valor = obtener_valor(sitio, campo)
        if valor:
            texto = str(valor)
            if texto not in partes:
                partes.append(texto)

    return " | ".join(partes)


def construir_fila(item):
    sitio = item.sitio

    fecha = item.fecha_planificada
    if isinstance(fecha, date):
        fecha = fecha.strftime("%d/%m/%Y")

    valores = [
        obtener_valor(sitio, "id_claro"),
        obtener_valor(sitio, "nombre"),
        obtener_valor(sitio, "region"),
        obtener_valor(sitio, "comuna"),
        obtener_valor(sitio, "tipo_zona"),
        obtener_valor(sitio, "direccion"),
        contacto_acceso(item),
        obtener_valor(sitio, "condiciones_acceso"),
        etiqueta_campo(item, "estado_permiso"),
        etiqueta_campo(item, "estado"),
        fecha or "",
        obtener_valor(item, "alerta_motor"),
    ]

    return [texto_seguro(valor) for valor in valores]


def coincide_filtros(item, filtros):
    sitio = item.sitio

    id_claro = normalizar(
        obtener_valor(sitio, "id_claro")
    )

    condicion = normalizar(
        obtener_valor(sitio, "condiciones_acceso")
    )

    permiso = str(item.estado_permiso or "")
    estado = str(item.estado or "")

    if filtros["id_claro"] not in id_claro:
        return False

    filtro_condicion = filtros["condicion"]

    if filtro_condicion == "__sin_informacion__":
        if condicion:
            return False
    elif filtro_condicion and filtro_condicion != condicion:
        return False

    if filtros["permiso"] and filtros["permiso"] != permiso:
        return False

    if filtros["estado"] and filtros["estado"] != estado:
        return False

    return True


@login_required
@require_GET
@rol_requerido("admin", "pm", "supervisor")
def exportar_asignacion_mensual(request, pk):

    planificacion = get_object_or_404(
        PlanificacionMensual,
        pk=pk,
    )

    formato = request.GET.get("formato", "xlsx").lower()

    if formato not in {"xlsx", "csv"}:
        return HttpResponse(
            "Formato de exportación no válido.",
            status=400,
        )

    filtros = {
        "id_claro": normalizar(
            request.GET.get("id_claro", "")
        ),
        "condicion": normalizar(
            request.GET.get("condicion_acceso", "")
        ),
        "permiso": request.GET.get("permiso", "").strip(),
        "estado": request.GET.get("estado", "").strip(),
    }

    queryset = (
        SitioPlanificado.objects
        .filter(
            planificacion=planificacion,
            activo_en_mes=True,
        )
        .select_related("sitio")
        .order_by(
            "sitio__region",
            "sitio__comuna",
            "sitio__id_claro",
        )
    )

    filas = [
        construir_fila(item)
        for item in queryset.iterator(chunk_size=500)
        if coincide_filtros(item, filtros)
    ]

    nombre = (
        f"planificacion_mensual_"
        f"{planificacion.anio}_"
        f"{planificacion.mes:02d}"
    )

    if formato == "csv":

        respuesta = HttpResponse(
            content_type="text/csv; charset=utf-8",
        )

        respuesta["Content-Disposition"] = (
            f'attachment; filename="{nombre}.csv"'
        )

        respuesta.write("\ufeff")

        escritor = csv.writer(respuesta)
        escritor.writerow(COLUMNAS)
        escritor.writerows(filas)

        return respuesta

    try:
        from openpyxl import Workbook
        from openpyxl.styles import (
            Alignment,
            Font,
            PatternFill,
        )
        from openpyxl.utils import get_column_letter
    except ImportError:
        return HttpResponse(
            "El módulo openpyxl no está instalado.",
            status=500,
        )

    libro = Workbook()
    hoja = libro.active
    hoja.title = "Planificación mensual"

    hoja.append(COLUMNAS)

    for fila in filas:
        hoja.append(fila)

    encabezado = PatternFill(
        fill_type="solid",
        fgColor="17365D",
    )

    for celda in hoja[1]:
        celda.fill = encabezado
        celda.font = Font(
            bold=True,
            color="FFFFFF",
        )
        celda.alignment = Alignment(
            vertical="center",
            wrap_text=True,
        )

    anchos = [
        17, 32, 23, 23, 19, 42,
        44, 31, 22, 22, 21, 48,
    ]

    for indice, ancho in enumerate(anchos, start=1):
        hoja.column_dimensions[
            get_column_letter(indice)
        ].width = ancho

    hoja.freeze_panes = "A2"
    hoja.auto_filter.ref = hoja.dimensions
    hoja.row_dimensions[1].height = 32

    buffer = io.BytesIO()
    libro.save(buffer)
    buffer.seek(0)

    respuesta = HttpResponse(
        buffer.getvalue(),
        content_type=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        ),
    )

    respuesta["Content-Disposition"] = (
        f'attachment; filename="{nombre}.xlsx"'
    )

    return respuesta
